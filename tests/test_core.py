"""Data model: normalisation, archive merging, table state, config parsing,
file handling. No display needed."""

import json
import os
from datetime import date

import pytest

from conftest import entry


# ── dates ───────────────────────────────────────────────────────────────────

def test_iso_weeks_across_the_year_boundary(tool):
    assert tool.iso_year_week_for_date(date(2026, 12, 31)) == (2026, 53)
    assert tool.iso_year_week_for_date(date(2027, 1, 3)) == (2026, 53)
    assert tool.iso_year_week_for_date(date(2027, 1, 4)) == (2027, 1)
    assert tool.week_after(2026, 53) == (2027, 1)
    assert tool.cw_range(2026, 53) == (date(2026, 12, 28), date(2027, 1, 1))   # Mon-Fri


@pytest.mark.parametrize("text", ["23/10/26", "23/10/2026", "2026-10-23", "23.10.2026"])
def test_safe_date_formats(tool, text):
    assert tool.safe_date(text) == date(2026, 10, 23)


# ── normalisation ───────────────────────────────────────────────────────────

def test_legacy_entry_fields(tool):
    a = tool.normalise_activity({"date": "12/10/2026", "user": "alice  a", "desc": "d",
                                 "task_no": 1, "task_title": "T1", "effort": "0,5", "task_wp": "wp06"})
    assert a["date"] == "2026-10-12"
    assert a["user"] == "Alice A"          # users.txt spelling
    assert a["description"] == "d"
    assert a["effort"] == 0.5
    assert a["task_wp"] == "WP6"


def test_missing_effort_is_one_day_null_is_zero(tool):
    assert tool.normalise_activity({"date": "2026-10-12", "user": "x"})["effort"] == 1
    assert tool.normalise_activity({"date": "2026-10-12", "user": "x", "effort": None})["effort"] == 0


def test_control_characters_removed(tool):
    a = tool.normalise_activity(entry("2026-10-12", "a\x0bb\x01c"))
    assert a["description"] == "a\nbc"


# ── archive merging ─────────────────────────────────────────────────────────

def test_merge_keeps_entries_that_were_not_loaded(tool):
    archived = [entry("2026-10-23", "d4")]
    table = [entry("2026-10-23", "forgot")]          # same day, user, task
    merged = tool.merge_week(archived, table, set())
    assert sorted(a["description"] for a in merged) == ["d4", "forgot"]


def test_merge_replaces_and_deletes_loaded_entries(tool):
    a, b, c = entry("2026-10-20", "a"), entry("2026-10-21", "b"), entry("2026-10-22", "c")
    loaded = {tool.loaded_key(x) for x in (a, b, c)}
    edited_b = dict(b, description="b2")
    merged = tool.merge_week([a, b, c], [a, edited_b], loaded)    # c deleted, b edited
    assert sorted(x["description"] for x in merged) == ["a", "b2"]


def test_merge_keeps_identical_entry_in_another_wp(tool):
    wp6 = entry("2026-10-20", "m", wp="WP6", title="Mgmt")
    slm = entry("2026-10-20", "m", wp="SLM", title="Mgmt")
    assert len(tool.merge_week([wp6], [slm], set())) == 2


def test_old_state_keys_are_honoured(tool):
    a = entry("2026-10-20", "a", user="Alice A")
    # 5-field stable key written by an older version, old user spelling
    tool.save_table_state({("20/10/2026", "alice  a", "Activity", "WP6", "1")})
    loaded = tool.load_table_state()
    assert tool._is_loaded(a, loaded)
    # 7-field key of the previous version
    assert tool._is_loaded(a, {tool.activity_key(a)})


# ── weekly archives ─────────────────────────────────────────────────────────

def test_week_archive_round_trip_and_monthly_dedup(tool):
    e1, e2 = entry("2026-10-20", "a"), entry("2026-10-20", "a")   # 2 x 0.5 PD, same text
    tool.write_week_archive(2026, 43, [e1, e2])
    data = tool.load_monthly_data_from_weekly_reports(2026, 10)
    assert len(data) == 2 and sum(x["effort"] for x in data) == 1.0


def test_onedrive_conflict_copy_is_not_counted(tool):
    tool.write_week_archive(2026, 43, [entry("2026-10-20", "a", effort=1)])
    wd = tool.share / "weekly_reports"
    (wd / "week_2026_CW43-LAPTOP.json").write_text(
        json.dumps([entry("2026-10-20", "a", effort=0.5)]), encoding="utf-8")
    data = tool.load_monthly_data_from_weekly_reports(2026, 10)
    assert [x["effort"] for x in data] == [1.0]
    assert any("LAPTOP" in s for s in tool.skipped_archives)


def test_unreadable_archive_is_reported_not_overwritten(tool):
    wd = tool.share / "weekly_reports"; wd.mkdir(exist_ok=True)
    (wd / "week_2026_CW43.json").write_text("", encoding="utf-8")
    assert tool.load_week_entries(2026, 43) == []
    assert "week_2026_CW43.json" in tool.skipped_archives
    with pytest.raises(RuntimeError):
        tool.write_week_archive(2026, 43, [entry("2026-10-20", "x")])


def test_atomic_write_retries_while_the_file_is_locked(tool, monkeypatch):
    calls = {"n": 0}
    real = os.replace

    def flaky(src, dst):
        calls["n"] += 1
        if calls["n"] < 3:
            raise PermissionError(13, "locked by OneDrive")
        return real(src, dst)
    monkeypatch.setattr(tool.os, "replace", flaky)
    monkeypatch.setattr(tool.time, "sleep", lambda s: None)
    target = tool.share / "x.json"
    tool.atomic_write_json(target, {"a": 1})
    assert json.loads(target.read_text()) == {"a": 1}
    assert not (tool.share / "x.json.tmp").exists()


def test_atomic_write_refuses_in_read_only(tool, monkeypatch):
    monkeypatch.setattr(tool, "read_only", True)
    with pytest.raises(PermissionError):
        tool.atomic_write_json(tool.share / "y.json", {})


# ── auto-archive at week change ─────────────────────────────────────────────

def test_week_change_archives_and_applies_pending_deletions(tool):
    keep, gone = entry("2026-10-20", "keep"), entry("2026-10-21", "gone")
    tool.write_week_archive(2026, 43, [keep, gone])
    cur = entry("2026-10-27", "cur-deleted")
    tool.write_week_archive(2026, 44, [cur])
    tool.write_json([keep])                                  # gone and cur deleted in table
    tool.save_table_state({tool.loaded_key(x) for x in (keep, gone, cur)})
    tool.Today.value = date(2026, 10, 28)
    data, reset, _ = tool.read_json()
    assert reset and data == []
    assert [a["description"] for a in tool.load_week_entries(2026, 43)] == ["keep"]
    assert tool.load_week_entries(2026, 44) == []


# ── config ──────────────────────────────────────────────────────────────────

def test_users_same_person_twice_keeps_first_spelling(tool):
    (tool.share / "config" / "users.txt").write_text("Fabio CIALFI\nfabio  cialfi\nBob B\n", encoding="utf-8")
    users, _ = tool.reload_config()
    assert users == ["Fabio CIALFI", "Bob B"]


def test_tasks_wp_spellings_and_no_wp(tool):
    (tool.share / "config" / "tasks.txt").write_text(
        "1 | A | P | wp6 | X\n2 | B | P | WP06 | X\n3 | C | P | WP-7 | X\n4 | D | Common | X\n",
        encoding="utf-8")
    _, tasks = tool.reload_config()
    assert sorted({wp for _, _, wp in tasks}) == ["", "WP6", "WP7"]


def test_config_files_with_bom_and_ansi(tool):
    (tool.share / "config" / "users.txt").write_bytes("﻿Nicolò R\n".encode("utf-8"))
    assert tool.load_users_from_config() == ["Nicolò R"]
    (tool.share / "config" / "users.txt").write_bytes("Nicolò R\n".encode("cp1252"))
    assert tool.load_users_from_config() == ["Nicolò R"]


def test_holiday_formats(tool):
    (tool.share / "config" / "holidays.txt").write_text(
        "2026-01-01 New Year\n01/05/2026 | Labour\n2026-12-25\tXmas\n2026-12-24 | Eve\nnonsense\n",
        encoding="utf-8")
    h = tool.load_holidays()
    assert h == {date(2026, 1, 1): "New Year", date(2026, 5, 1): "Labour",
                 date(2026, 12, 25): "Xmas", date(2026, 12, 24): "Eve"}


# ── import / export ─────────────────────────────────────────────────────────

def test_own_csv_export_reimports(tool, tmp_path):
    import csv
    p = tmp_path / "x.csv"
    row = entry("2026-10-20", 'a, b, c, d, e, f, g, h, i\nline2; "q"')
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["date", "user", "entry_type", "task_wp", "task_no",
                                          "task_title", "effort", "description"], delimiter=";")
        w.writeheader(); w.writerow(row)
    ok, rejected = tool.import_activity_table(p)
    assert rejected == [] and ok[0]["description"] == row["description"]


def test_csv_saved_by_excel_in_ansi(tool, tmp_path):
    p = tmp_path / "x.csv"
    p.write_bytes("date;user;entry_type;task_no;task_title;effort;description\n"
                  "20/10/2026;Alice A;Activity;1;T1;0,5;caffè\n".encode("cp1252"))
    ok, rejected = tool.import_activity_table(p)
    assert rejected == [] and ok[0]["description"] == "caffè" and ok[0]["effort"] == 0.5


def test_xlsx_export_text_not_formula_and_total_activities_only(tool, tmp_path):
    import openpyxl
    out = tmp_path / "x.xlsx"
    tool.export_to_xlsx([entry("2026-10-20", "=> follow-up", effort=1),
                         entry("2026-10-21", "", etype="Holiday", effort=1)], str(out))
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(2, 7).data_type == "s"
    assert ws.cell(4, 6).value == 1   # TOTAL row: leave excluded


# ── version check ───────────────────────────────────────────────────────────

def test_version_parsing(tool):
    assert tool._parse_version("1.2.10") > tool._parse_version("1.2.9")


def test_version_file_with_bom(T, tool, monkeypatch, tmp_path):
    rel = tmp_path / "rel"; rel.mkdir()
    (rel / "VERSION.txt").write_bytes("﻿99.0.0\r\n".encode("utf-8"))
    monkeypatch.setattr(T, "RELEASE_DIR", str(rel))
    monkeypatch.undo()   # check_version was stubbed by the fixture: use the real one
    monkeypatch.setattr(T, "RELEASE_DIR", str(rel))
    import importlib
    real = importlib.import_module("Exostaff_Activity_Report_Tool").__dict__["check_version"]
    latest, update = real()
    assert latest == "99.0.0" and update
