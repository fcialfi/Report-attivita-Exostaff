"""Reports: monthly Word (real template), weekly summary, task trend,
next-week plan. No display needed."""

import shutil

import docx
import pytest

from conftest import REPO, entry

TEMPLATE = next((REPO / "config").glob("*.docx"))


@pytest.fixture
def real_config(tool):
    for name in ("users.txt", "tasks.txt", "holidays.txt"):
        shutil.copy(REPO / "config" / name, tool.share / "config" / name)
    tool.reload_config()
    tool.reload_holidays()
    return tool


def user_table(doc, user):
    for u, table in [(u, t) for u, t in __import__("Exostaff_Activity_Report_Tool").get_user_tables(doc)]:
        if u == user:
            return table
    raise AssertionError(f"no table for {user}")


def rows_with_effort(table):
    return {int(r.cells[0].text.strip()): float(r.cells[2].text)
            for r in table.rows if r.cells[0].text.strip().isdigit() and r.cells[2].text.strip()}


def summary(doc, title):
    for t in doc.tables:
        if title in " ".join(c.text for c in t.rows[0].cells):
            return {r.cells[0].text.strip(): r.cells[2].text.strip() for r in t.rows[1:]}
    raise AssertionError(title)


def make(tool, data, tmp_path, y=2026, m=10):
    out = tmp_path / "report.docx"
    warnings = tool.generate_monthly_docx(TEMPLATE, data, y, m, str(out))
    return docx.Document(str(out)), warnings


def test_monthly_totals_match_entries(real_config, tmp_path):
    T = real_config
    title9 = T._task_titles[(9, "")]
    data = [entry("2026-10-05", "satsim a", task=9, wp="", title=title9, effort=1, user="Fabio CIALFI"),
            entry("2026-10-06", "satsim b", task=9, wp="", title=title9, effort=0.5, user="Fabio CIALFI"),
            entry("2026-10-07", "", etype="Holiday", effort=1, user="Fabio CIALFI"),
            entry("2026-11-02", "next month", task=9, wp="", title=title9, effort=1, user="Fabio CIALFI")]
    doc, warnings = make(T, data, tmp_path)
    assert warnings == []
    assert sum(rows_with_effort(user_table(doc, "Fabio CIALFI")).values()) == 1.5
    eff = summary(doc, "Total effort")
    absn = summary(doc, "Total absence")
    assert eff["FaC"] == "1.5" and eff["TOTAL"] == "1.5"
    assert absn["FaC"] == "1" and absn["TOTAL"] == "1"


def test_monthly_month_text_updated(real_config, tmp_path):
    doc, _ = make(real_config, [], tmp_path, 2027, 1)
    xml = "".join(p._p.xml for p in doc.paragraphs)
    for h in doc.sections[0].header.paragraphs:
        xml += h._p.xml
    assert "April 2026" not in xml and "04-2026" not in xml


def test_monthly_rows_add_up_to_totals(real_config, tmp_path):
    T = real_config
    t9, t15 = T._task_titles[(9, "")], T._task_titles[(15, "")]
    data = ([entry("2026-10-05", f"a{i}", task=9, wp="", title=t9, effort=0.333, user="Fabio CIALFI")
             for i in range(3)]
            + [entry("2026-10-06", f"b{i}", task=15, wp="", title=t15, effort=0.125, user="Fabio CIALFI")
               for i in range(3)])
    doc, _ = make(T, data, tmp_path)
    rows = rows_with_effort(user_table(doc, "Fabio CIALFI"))
    assert round(sum(rows.values()), 2) == float(summary(doc, "Total effort")["FaC"])


def test_monthly_legacy_entries_split_by_recovered_wp(tool, tmp_path):
    T = tool
    shutil.copy(REPO / "config" / "users.txt", T.share / "config" / "users.txt")
    T.reload_config()
    # T1 (WP6) and S1 (SLM) share task number 1; old entries carry no WP
    data = [entry("2026-10-05", "wp6 work", task=1, wp="", title="T1", effort=3, user="Fabio CIALFI"),
            entry("2026-10-06", "slm work", task=1, wp="", title="S1", effort=1, user="Fabio CIALFI")]
    groups = {}
    for a in data:
        groups.setdefault(T.entry_wp(a), []).append(a)
    assert sorted(groups) == ["SLM", "WP6"]


def test_weekly_summary_plain_lines_and_plan_by_wp(tool):
    data = [entry("2026-10-20", "did a", task=1, wp="WP6"),
            entry("2026-10-21", "did b", task=3, wp="WP7"),
            entry("2026-10-21", "slm", task=1, wp="SLM")]
    plan = {"P": ["older plan line"], "WP7": ["x"]}      # "P" = project of tasks.txt
    txt = tool.generate_weekly_text(data, "2026-10-19", "2026-10-23", cw=43, year=2026,
                                    next_week_plan=plan)
    assert "o\t" not in txt
    assert "WP6:\nDid a." in txt or "WP6:\ndid a." in txt
    nxt = txt.split("Next week", 1)[1]
    assert "Older plan line." in nxt or "older plan line." in nxt
    assert "P:" not in nxt                                # no project section


def test_regroup_plan_never_to_slm(tool):
    out = tool.regroup_plan({"P": ["a"]}, ["SLM", "WP6", "WP7"])
    assert out == {"WP6": ["a"]}


def test_task_trend_and_export(tool, tmp_path):
    import openpyxl
    data = [entry("2026-09-10", "a", task=1, wp="WP6", effort=1),
            entry("2026-11-10", "b", task=1, wp="WP6", effort=0.5),
            entry("2026-11-11", "c", task=3, wp="WP7", effort=2, user="bob  b"),
            entry("2026-11-12", "s", task=1, wp="SLM", effort=1),
            entry("2026-11-13", "z", task=3, wp="WP7", effort=0)]
    months = tool.month_span((2026, 7), (2026, 12))
    rows = tool.task_trend(data, months)
    assert [(r["wp"], r["task_no"], r["total"]) for r in rows] == [("WP6", 1, 1.5), ("WP7", 3, 2.0)]
    assert tool.trim_empty_months(rows, months) == [(2026, 9), (2026, 10), (2026, 11)]
    assert [r["total"] for r in tool.task_trend(data, months, user="Bob B")] == [2.0]
    out = tmp_path / "trend.xlsx"
    tool.export_trend_xlsx(rows, tool.trim_empty_months(rows, months), str(out), "t")
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(ws.max_row, 1).value == "TOTAL" and ws.cell(ws.max_row, 6).value == 3.5
    assert len(ws._charts) == 1
